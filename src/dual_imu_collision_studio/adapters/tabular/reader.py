from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook

from ...domain.data_workspace.model import (
    ColumnMapping,
    HeaderProposal,
    ImportedSamples,
)


class TabularReader:
    @staticmethod
    def propose_header(preview_rows: tuple[tuple[Any, ...], ...]) -> HeaderProposal:
        candidates: list[tuple[int, int, tuple[str, ...]]] = []
        for index, row in enumerate(preview_rows):
            columns = tuple(
                str(value).strip()
                for value in row
                if value is not None and str(value).strip()
            )
            if len(columns) < 3 or len(set(columns)) != len(columns):
                continue
            lowered = tuple(value.lower() for value in columns)
            score = sum("internal" in value and "accel" in value for value in lowered)
            score += sum(
                ("diff" in value or "external" in value) and "accel" in value
                for value in lowered
            )
            score += sum(value.endswith((" x", " y", " z")) for value in lowered)
            candidates.append((score, -index, columns))
        if not candidates:
            raise ValueError(
                "no header row with at least three unique columns was detected"
            )
        _, negative_index, columns = max(candidates)
        header_row = -negative_index + 1
        return HeaderProposal(header_row, header_row + 1, columns)

    def worksheets(self, path: str | Path) -> tuple[str, ...]:
        source = Path(path)
        if source.suffix.lower() != ".xlsx":
            return ()
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            return tuple(workbook.sheetnames)
        finally:
            workbook.close()

    def preview(
        self, path: str | Path, worksheet: str | None = None, limit: int = 12
    ) -> tuple[tuple[Any, ...], ...]:
        return tuple(self._rows(Path(path), worksheet)[:limit])

    def read(
        self,
        path: str | Path,
        mapping: ColumnMapping,
        worksheet: str | None = None,
    ) -> ImportedSamples:
        rows = self._rows(Path(path), worksheet)
        if mapping.data_start_row > len(rows):
            raise ValueError("Data Start Row exceeds the source row count")
        headers = self._headers(rows, mapping)
        indexes = {name: index for index, name in enumerate(headers)}
        requested = mapping.internal_columns + mapping.external_columns
        if mapping.time_column is not None:
            requested += (mapping.time_column,)
        missing = [name for name in requested if name not in indexes]
        if missing:
            raise ValueError("mapped columns do not exist: " + ", ".join(missing))
        data_rows = rows[mapping.data_start_row - 1 :]
        if not data_rows:
            raise ValueError("recording contains no data rows")
        data_rows, warnings = self._trim_trailing_incomplete_rows(
            data_rows, requested, indexes, mapping.data_start_row
        )
        internal = (
            self._numeric_matrix(
                data_rows, mapping.internal_columns, indexes, mapping.data_start_row
            )
            if mapping.internal_columns
            else np.zeros((len(data_rows), 3))
        )
        external = (
            self._numeric_matrix(
                data_rows, mapping.external_columns, indexes, mapping.data_start_row
            )
            if mapping.external_columns
            else np.zeros((len(data_rows), 3))
        )
        interpolated_indices: tuple[int, ...] = ()
        if mapping.time_column is None:
            rate = float(mapping.sample_rate_hz)
            time_s = np.arange(len(data_rows), dtype=np.float64) / rate
            discontinuity = np.zeros(len(data_rows), dtype=np.bool_)
        else:
            raw_time = self._numeric_column(
                data_rows, mapping.time_column, indexes, mapping.data_start_row
            )
            factor = {"s": 1.0, "ms": 1e-3, "us": 1e-6, "µs": 1e-6}[mapping.time_unit]
            time_s = raw_time * factor
            rate, discontinuity = self._validate_time(time_s)
            time_s, internal, external, interpolated_indices = self._interpolate_gaps(
                time_s, internal, external, rate, discontinuity
            )
            discontinuity = np.zeros(len(time_s), dtype=np.bool_)
            if interpolated_indices:
                warnings += (
                    f"{len(interpolated_indices)} sample(s) linearly interpolated across timestamp gaps",
                )
        return ImportedSamples(
            time_s,
            internal if mapping.internal_columns else None,
            external if mapping.external_columns else None,
            rate,
            discontinuity,
            warnings,
            interpolated_indices,
        )

    @staticmethod
    def _interpolate_gaps(time_s, internal, external, rate, discontinuity):
        gaps = np.flatnonzero(discontinuity)
        if not len(gaps):
            return time_s, internal, external, ()
        period = 1.0 / rate
        steps = np.ones(len(time_s), dtype=np.int64)
        for right in gaps:
            elapsed = float(time_s[right] - time_s[right - 1])
            ratio = elapsed / period
            if not np.isfinite(ratio):
                raise ValueError(
                    "timestamp gap cannot be represented at the inferred Sample Rate"
                )
            intervals = round(ratio)
            missing = intervals - 1
            if missing >= 4:
                raise ValueError(
                    f"recording rejected: {missing} consecutive missing samples between "
                    f"{time_s[right - 1]:.12g} and {time_s[right]:.12g} s (maximum 3)"
                )
            if missing < 1 or abs(elapsed / intervals - period) / period > 0.02 + 1e-12:
                raise ValueError(
                    f"recording rejected: timestamp gap before sample {right} exceeds ±2% alignment"
                )
            steps[right] = intervals
        measured = np.cumsum(steps) - 1
        count = int(measured[-1]) + 1
        times = np.empty(count, dtype=np.float64)
        inside = np.empty((count, 3), dtype=np.float64)
        outside = np.empty((count, 3), dtype=np.float64)
        times[measured], inside[measured], outside[measured] = (
            time_s,
            internal,
            external,
        )
        inserted: list[int] = []
        for right in gaps:
            left = right - 1
            for offset in range(1, int(steps[right])):
                index = int(measured[left]) + offset
                fraction = offset / int(steps[right])
                times[index] = time_s[left] + fraction * (time_s[right] - time_s[left])
                inside[index] = (1 - fraction) * internal[left] + fraction * internal[
                    right
                ]
                outside[index] = (1 - fraction) * external[left] + fraction * external[
                    right
                ]
                inserted.append(index)
        if (
            not np.all(np.diff(times) > 0)
            or not np.all(np.isfinite(inside))
            or not np.all(np.isfinite(outside))
        ):
            raise ValueError(
                "interpolation produced invalid timestamps or non-finite IMU values"
            )
        return times, inside, outside, tuple(inserted)

    @staticmethod
    def _trim_trailing_incomplete_rows(
        rows: list[tuple[Any, ...]],
        columns: tuple[str, ...],
        indexes: dict[str, int],
        first_row: int,
    ) -> tuple[list[tuple[Any, ...]], tuple[str, ...]]:
        complete: list[bool] = []
        for row in rows:
            try:
                values = [float(row[indexes[name]]) for name in columns]
            except (IndexError, TypeError, ValueError):
                complete.append(False)
                continue
            complete.append(all(np.isfinite(value) for value in values))
        try:
            first_incomplete = complete.index(False)
        except ValueError:
            return rows, ()
        if any(complete[first_incomplete + 1 :]):
            source_row = first_row + first_incomplete
            raise ValueError(
                f"mapped value gap inside recording at source row {source_row}"
            )
        if first_incomplete == 0:
            raise ValueError("recording contains no complete mapped data rows")
        omitted = len(rows) - first_incomplete
        warning = f"{omitted} trailing source row(s) omitted because mapped values were incomplete"
        return rows[:first_incomplete], (warning,)

    @staticmethod
    def _rows(path: Path, worksheet: str | None) -> list[tuple[Any, ...]]:
        suffix = path.suffix.lower()
        if suffix == ".csv":
            with path.open("r", newline="", encoding="utf-8-sig") as stream:
                return [tuple(row) for row in csv.reader(stream)]
        if suffix == ".xlsx":
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                selected = worksheet or workbook.sheetnames[0]
                if selected not in workbook.sheetnames:
                    raise ValueError(f"worksheet does not exist: {selected}")
                return [
                    tuple(row) for row in workbook[selected].iter_rows(values_only=True)
                ]
            finally:
                workbook.close()
        raise ValueError("only CSV and .xlsx recordings are supported")

    @staticmethod
    def _headers(
        rows: list[tuple[Any, ...]], mapping: ColumnMapping
    ) -> tuple[str, ...]:
        if mapping.header_row is None:
            width = max(len(row) for row in rows)
            return tuple(f"Column {index + 1}" for index in range(width))
        if mapping.header_row < 1 or mapping.header_row > len(rows):
            raise ValueError("Header Row is outside the source")
        source_headers = list(rows[mapping.header_row - 1])
        while source_headers and (
            source_headers[-1] is None or not str(source_headers[-1]).strip()
        ):
            source_headers.pop()
        headers = tuple(
            str(value).strip() if value is not None else "" for value in source_headers
        )
        if any(not name for name in headers) or len(set(headers)) != len(headers):
            raise ValueError("Header Row must contain unique non-empty column names")
        return headers

    @classmethod
    def _numeric_matrix(
        cls,
        rows: list[tuple[Any, ...]],
        columns: tuple[str, str, str],
        indexes: dict[str, int],
        first_row: int,
    ) -> np.ndarray:
        return np.column_stack(
            [cls._numeric_column(rows, name, indexes, first_row) for name in columns]
        )

    @staticmethod
    def _numeric_column(
        rows: list[tuple[Any, ...]], name: str, indexes: dict[str, int], first_row: int
    ) -> np.ndarray:
        values: list[float] = []
        index = indexes[name]
        for offset, row in enumerate(rows):
            try:
                value = float(row[index])
            except (IndexError, TypeError, ValueError):
                raise ValueError(
                    f"non-numeric value in {name} at source row {first_row + offset}"
                ) from None
            if not np.isfinite(value):
                raise ValueError(
                    f"non-finite value in {name} at source row {first_row + offset}"
                )
            values.append(value)
        return np.asarray(values, dtype=np.float64)

    @staticmethod
    def _validate_time(time_s: np.ndarray) -> tuple[float, np.ndarray]:
        if len(time_s) < 2:
            raise ValueError("a Time column requires at least two data rows")
        delta = np.diff(time_s)
        if (
            not np.all(np.isfinite(time_s))
            or not np.all(np.isfinite(delta))
            or np.any(delta <= 0)
        ):
            raise ValueError("Time column must be finite and strictly increasing")
        base_dt = float(np.median(delta))
        discontinuity_delta = delta > 1.5 * base_dt
        ordinary = delta[~discontinuity_delta]
        if ordinary.size and np.any(np.abs(ordinary - base_dt) / base_dt > 0.02):
            raise ValueError("ordinary timestamp jitter exceeds ±2%")
        discontinuity = np.zeros(len(time_s), dtype=np.bool_)
        discontinuity[1:] = discontinuity_delta
        return 1.0 / base_dt, discontinuity
